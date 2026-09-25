"""Чтение xlsx/ods/csv/json/jsonl: лимиты, utf-8-sig, корректные ошибки."""
import csv
import json
import logging
from pathlib import Path
from typing import Any

import openpyxl

logger = logging.getLogger(__name__)

MAX_PREVIEW_BYTES = 200 * 1024 * 1024

# Ошибки вычисления формул (EN + RU): закэшированное значение ячейки.
# Показывать их в ревью как есть — мусор («REP!» из жалобы); втягиваем
# пустыми, а количество отдаём в stats для предпроверки импорта.
# Совпадение строго целой ячейкой — «#1» и прочие тексты не трогаем.
EXCEL_ERRORS = frozenset({
    "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!",
    "#GETTING_DATA", "#SPILL!", "#CALC!", "#FIELD!", "#CONNECT!",
    "#BLOCKED!", "#UNKNOWN!",
    "#ДЕЛ/0!", "#Н/Д", "#ИМЯ?", "#ПУСТО!", "#ЧИСЛО!", "#ССЫЛКА!",
    "#ЗНАЧ!", "#ПОЛУЧЕНИЕ_ДАННЫХ",
})


def is_excel_error(value) -> bool:
    """Ячейка — код ошибки формулы (а не данные)."""
    try:
        return isinstance(value, str) and value.strip().upper() in EXCEL_ERRORS
    except Exception:
        return False


def _clean_cell(value, stats: dict | None):
    """Ошибка формулы -> '' (+счётчик), остальное как есть."""
    if is_excel_error(value):
        if stats is not None:
            try:
                stats["formula_errors"] = int(stats.get("formula_errors", 0)) + 1
            except Exception:
                pass
        return ""
    return value


def _check_size(file_path: str) -> None:
    size = Path(file_path).stat().st_size
    if size > MAX_PREVIEW_BYTES:
        raise ValueError("Файл слишком большой для превью (>200 МБ)")


def _validate_delimiter(delimiter: str) -> None:
    if not isinstance(delimiter, str) or len(delimiter) != 1:
        raise ValueError("Разделитель CSV — ровно один символ")


class FileReader:
    """Читает различные форматы файлов и возвращает данные."""

    @staticmethod
    def detect_file_type(file_path: str) -> str:
        ext = Path(file_path).suffix.lower()
        if ext == ".xlsx":
            return "excel"
        if ext == ".xls":
            raise ValueError(".xls не поддерживается openpyxl — сохраните как .xlsx")
        if ext == ".ods":
            return "ods"
        if ext == ".csv":
            return "csv"
        if ext == ".json":
            return "json"
        if ext == ".jsonl":
            return "jsonl"
        return "unknown"

    @staticmethod
    def read_excel_sheets(file_path: str) -> list:
        _check_size(file_path)
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        try:
            return list(wb.sheetnames)
        except Exception as e:
            raise ValueError(f"Не удалось открыть Excel-файл: {e}") from e
        finally:
            wb.close()

    @staticmethod
    def _ods_table(file_path: str, sheet_name: str | None):
        """Открывает ODS и возвращает (doc, table). Нужен odfpy."""
        try:
            from odf.opendocument import load as _ods_load
        except ImportError as e:
            raise ValueError("Для .ods установите odfpy: pip install odfpy") from e
        _check_size(file_path)
        try:
            doc = _ods_load(file_path)
        except Exception as e:
            raise ValueError(f"Не удалось открыть ODS-файл: {e}") from e
        from odf.table import Table as _Table
        tables = doc.getElementsByType(_Table)
        if not tables:
            raise ValueError("В ODS нет листов")
        if sheet_name:
            for t in tables:
                if t.getAttribute("name") == sheet_name:
                    return doc, t
            raise ValueError(f"Лист {sheet_name!r} не найден")
        return doc, tables[0]

    @staticmethod
    def _ods_rows(table, max_rows: int = 0, stats: dict | None = None) -> list:
        """Строки листа как списки строк (с учётом repeated)."""
        from odf.table import TableRow as _Row, TableCell as _Cell
        from odf.text import P as _P
        out = []
        for row in table.getElementsByType(_Row):
            if max_rows and len(out) >= max_rows:
                break
            cells = []
            for cell in row.getElementsByType(_Cell):
                repeat = cell.getAttribute("numbercolumnsrepeated")
                try:
                    repeat = int(repeat) if repeat else 1
                except (TypeError, ValueError):
                    repeat = 1
                texts = []
                for p in cell.getElementsByType(_P):
                    parts = []
                    for node in p.childNodes:
                        if node.nodeType == node.TEXT_NODE:
                            parts.append(node.data)
                    texts.append("".join(parts).strip())
                value = "\n".join(t for t in texts if t)
                value = _clean_cell(value, stats)
                cells.extend([value] * min(repeat, 256))
            out.append(cells)
        return out

    @staticmethod
    def read_ods_sheets(file_path: str) -> list:
        from odf.table import Table as _Table
        try:
            from odf.opendocument import load as _ods_load
        except ImportError as e:
            raise ValueError("Для .ods установите odfpy: pip install odfpy") from e
        _check_size(file_path)
        try:
            doc = _ods_load(file_path)
            return [t.getAttribute("name") or f"Лист{i + 1}"
                    for i, t in enumerate(doc.getElementsByType(_Table))]
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"Не удалось открыть ODS-файл: {e}") from e

    @staticmethod
    def read_ods_preview(file_path: str, sheet_name: str, max_rows: int = 100) -> dict:
        try:
            _doc, table = FileReader._ods_table(file_path, sheet_name)
            rows = FileReader._ods_rows(table, max_rows)
            if not rows:
                return {"headers": [], "rows": []}
            return {"headers": rows[0], "rows": rows[1:]}
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"Не удалось прочитать ODS-файл: {e}") from e

    @staticmethod
    def read_ods_data(file_path: str, sheet_name: str, header_row: int = 0,
                       stats: dict | None = None) -> list:
        try:
            _doc, table = FileReader._ods_table(file_path, sheet_name)
            rows = FileReader._ods_rows(table, 0, stats)
            if not rows or header_row >= len(rows):
                return []
            headers = [h if h else f"col_{i}" for i, h in enumerate(rows[header_row])]
            data = []
            for row in rows[header_row + 1:]:
                if all((c or "").strip() == "" for c in row):
                    continue
                data.append({h: _clean_cell(row[i] if i < len(row) else "", stats)
                             for i, h in enumerate(headers)})
            return data
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"Ошибка чтения ODS: {e}") from e

    @staticmethod
    def read_excel_preview(file_path: str, sheet_name: str, max_rows: int = 100) -> dict:
        _check_size(file_path)
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        try:
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"Лист {sheet_name!r} не найден")
            ws = wb[sheet_name]
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= max_rows:
                    break
                # Превью маппинга — тоже без мусора формул (счётчик не нужен:
                # он в data-пути и показывается в предпроверке импорта).
                rows.append([("" if cell is None else str(_clean_cell(cell, None)))
                             for cell in row])
            if not rows:
                return {"headers": [], "rows": []}
            return {"headers": rows[0], "rows": rows[1:]}
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"Не удалось прочитать Excel-файл: {e}") from e
        finally:
            wb.close()

    @staticmethod
    def read_csv_preview(
        file_path: str, encoding: str = "utf-8-sig", delimiter: str = ",", max_rows: int = 100
    ) -> dict:
        _validate_delimiter(delimiter)
        _check_size(file_path)
        try:
            rows = []
            with open(file_path, encoding=encoding, newline="") as f:
                reader = csv.reader(f, delimiter=delimiter)
                for i, row in enumerate(reader):
                    if i >= max_rows:
                        break
                    rows.append(row)
            if not rows:
                return {"headers": [], "rows": []}
            return {"headers": rows[0], "rows": rows[1:]}
        except UnicodeDecodeError as e:
            raise ValueError("Не удалось определить кодировку. Попробуйте выбрать другую.") from e
        except Exception as e:
            raise ValueError(f"Не удалось прочитать CSV-файл: {e}") from e

    @staticmethod
    def _json_headers_and_rows(data: list, max_rows: int) -> dict:
        items = [x for x in data if isinstance(x, dict)]
        if not items:
            return {"headers": [], "rows": []}
        headers = list(items[0].keys())
        rows = [[str(it.get(h, "")) for h in headers] for it in items[:max_rows]]
        return {"headers": headers, "rows": rows}

    @staticmethod
    def read_json_preview(file_path: str, max_rows: int = 100) -> dict:
        _check_size(file_path)
        try:
            with open(file_path, encoding="utf-8-sig") as f:
                data = json.load(f)
            if not isinstance(data, list):
                raise ValueError("JSON должен содержать массив объектов")
            return FileReader._json_headers_and_rows(data, max_rows)
        except json.JSONDecodeError as e:
            raise ValueError(f"Ошибка парсинга JSON: {e}") from e

    @staticmethod
    def read_jsonl_preview(file_path: str, max_rows: int = 100) -> dict:
        _check_size(file_path)
        headers: list = []
        rows: list = []
        errors: list = []
        try:
            with open(file_path, encoding="utf-8-sig") as f:
                for i, line in enumerate(f):
                    if len(rows) >= max_rows:
                        break
                    s = line.strip()
                    if not s:
                        continue
                    try:
                        item = json.loads(s)
                    except json.JSONDecodeError as e:
                        errors.append({"line": i + 1, "error": str(e)})
                        continue
                    if not isinstance(item, dict):
                        errors.append({"line": i + 1, "error": "not an object"})
                        continue
                    if not headers:
                        headers = list(item.keys())
                    rows.append([str(item.get(h, "")) for h in headers])
            result: dict[str, Any] = {"headers": headers, "rows": rows}
            if errors:
                result["errors"] = errors
            return result
        except Exception as e:
            raise ValueError(f"Не удалось прочитать JSONL-файл: {e}") from e

    @staticmethod
    def read_excel_data(file_path: str, sheet_name: str, header_row: int = 0,
                        stats: dict | None = None) -> list:
        _check_size(file_path)
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        try:
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"Лист {sheet_name!r} не найден")
            ws = wb[sheet_name]
            data = []
            headers: list = []
            seen_header = False
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i < header_row:
                    continue
                if not seen_header:
                    headers = [str(h) if h not in (None, "") else f"col_{j}"
                               for j, h in enumerate(row)]
                    seen_header = True
                    continue
                if all(c is None or str(c).strip() == "" for c in row):
                    continue
                try:
                    from importer import _cell_str as _norm
                except Exception:
                    _norm = None
                row_dict = {}
                for j, header in enumerate(headers):
                    v = row[j] if j < len(row) else None
                    if v is None:
                        row_dict[header] = ""
                        continue
                    v = _clean_cell(v, stats)
                    if v == "":
                        row_dict[header] = ""
                        continue
                    try:
                        row_dict[header] = _norm(v) if _norm else v
                    except Exception:
                        row_dict[header] = v
                data.append(row_dict)
            return data
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"Ошибка чтения Excel: {e}") from e
        finally:
            wb.close()

    @staticmethod
    def read_csv_data(
        file_path: str, encoding: str = "utf-8-sig", delimiter: str = ",", header_row: int = 0
    ) -> list:
        _validate_delimiter(delimiter)
        _check_size(file_path)
        try:
            with open(file_path, encoding=encoding, newline="") as f:
                rows = list(csv.reader(f, delimiter=delimiter))
            if not rows or header_row >= len(rows):
                return []
            headers = [h if h else f"col_{i}" for i, h in enumerate(rows[header_row])]
            data = []
            for row in rows[header_row + 1:]:
                if all((c or "").strip() == "" for c in row):
                    continue
                data.append({h: (row[i] if i < len(row) else "") for i, h in enumerate(headers)})
            return data
        except UnicodeDecodeError as e:
            raise ValueError("Ошибка кодировки") from e
        except Exception as e:
            raise ValueError(f"Ошибка чтения CSV: {e}") from e

    @staticmethod
    def read_json_data(file_path: str) -> tuple:
        """Возвращает (rows, errors): битые/не-объекты не теряются молча."""
        _check_size(file_path)
        try:
            with open(file_path, encoding="utf-8-sig") as f:
                data = json.load(f)
            if not isinstance(data, list):
                raise ValueError("JSON должен содержать массив объектов")
            result, errors = [], []
            for i, item in enumerate(data):
                if not isinstance(item, dict):
                    errors.append({"line": i + 1, "error": "not an object, skipped"})
                    continue
                if not item:
                    continue
                row_dict = {}
                for key, value in item.items():
                    if isinstance(value, (dict, list)):
                        row_dict[key] = json.dumps(value, ensure_ascii=False)
                    else:
                        row_dict[key] = "" if value is None else value
                result.append(row_dict)
            return result, errors
        except json.JSONDecodeError as e:
            raise ValueError(f"Ошибка парсинга JSON: {e}") from e

    @staticmethod
    def read_jsonl_data(file_path: str) -> tuple:
        _check_size(file_path)
        result, errors = [], []
        try:
            with open(file_path, encoding="utf-8-sig") as f:
                for line_num, line in enumerate(f, 1):
                    s = line.strip()
                    if not s:
                        continue
                    try:
                        item = json.loads(s)
                    except json.JSONDecodeError as e:
                        errors.append({"line": line_num, "error": str(e)})
                        continue
                    if not isinstance(item, dict):
                        errors.append({"line": line_num, "error": "not an object, skipped"})
                        continue
                    row_dict = {}
                    for key, value in item.items():
                        if isinstance(value, (dict, list)):
                            row_dict[key] = json.dumps(value, ensure_ascii=False)
                        else:
                            row_dict[key] = "" if value is None else value
                    result.append(row_dict)
            return result, errors
        except Exception as e:
            raise ValueError(f"Ошибка чтения JSONL: {e}") from e
