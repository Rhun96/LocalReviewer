"""Универсальная аналитика поверх существующих данных (ТЗ Analytics v1.0).

Принципы:
- никаких новых таблиц/миграций: cases/annotations/case_errors/bugs/history;
- цифра на карточке считается ТЕМ ЖЕ предикатом, что drill-down
  (count_filtered_cases + drill_filters) — ключевой тест ТЗ держится
  конструкцией, а не надеждой;
- нет данных — «Нет данных», а не ноль (функции возвращают None-маркеры);
- тяжёлое — агрегатами SQL, не вытягиванием кейсов;
- универсальность: статусы — через base-семантику профиля (свои коды
  складываются), категории/тяжести — только существующие в проекте.
"""
import logging

from database import db

logger = logging.getLogger(__name__)

BASES = ("good", "bad", "uncertain", "skip", "duplicate")


def _codes_by_base(project_path: str) -> dict:
    """base -> [коды профиля]. Свои коды поддерживаются автоматически."""
    try:
        from review_profile_service import code_to_base
        mapping = code_to_base(project_path) or {}
    except Exception:
        mapping = {}
    out = {b: [] for b in BASES}
    for code, base in mapping.items():
        if base in out and code not in out[base]:
            out[base].append(code)
    if not any(out.values()):
        from constants import STATUSES as _ST
        for code in _ST:
            if code == "unreviewed":
                continue
            if code in BASES and code not in out[code]:
                out[code].append(code)
    return out


def scope_filters(scope: dict | None) -> dict:
    """Фильтры скоупа сводки: файл + период (без статусов)."""
    scope = scope or {}
    out: dict = {}
    if scope.get("file_id"):
        try:
            out["file_id"] = int(scope["file_id"])
        except (TypeError, ValueError):
            pass
    for k in ("reviewed_from", "reviewed_to"):
        v = str(scope.get(k, "") or "").strip()
        if v:
            out[k] = v[:10]
    return out


def reviewed_codes(project_path: str) -> list:
    codes = []
    for base in BASES:
        codes.extend(_codes_by_base(project_path).get(base, []))
    return codes


def card_counts(project_path: str, scope: dict | None) -> dict:
    """Карточки сводки + готовые drill-фильтры (число == drill-down)."""
    from filter_service import count_filtered_cases
    base = scope_filters(scope)
    rev = reviewed_codes(project_path)
    by_base = _codes_by_base(project_path)
    cards = {}
    total = count_filtered_cases(project_path, dict(base))
    cards["total"] = {"value": total, "filters": dict(base)}
    rflt = dict(base, statuses=list(rev))
    reviewed = count_filtered_cases(project_path, rflt)
    cards["reviewed"] = {"value": reviewed, "filters": rflt}
    for b in BASES:
        codes = by_base.get(b, [])
        flt = dict(base, statuses=list(codes))
        cards[b] = {"value": count_filtered_cases(project_path, flt)
                    if codes else 0, "filters": flt}
    unrev = total - reviewed
    cards["unreviewed"] = {"value": unrev,
                           "filters": dict(base, statuses=["unreviewed"])}
    return {"cards": cards, "scope": base}


def _pct(part: int, whole: int):
    if not whole:
        return None
    return round(100.0 * part / whole, 1)


def percentages(card_counts_result: dict) -> dict:
    """Good/Bad/неопределённые доли. None — «Нет данных», не ноль."""
    cards = (card_counts_result or {}).get("cards", {})
    rev = (cards.get("reviewed") or {}).get("value", 0) or 0
    out = {
        "good": _pct((cards.get("good") or {}).get("value", 0), rev),
        "bad": _pct((cards.get("bad") or {}).get("value", 0), rev),
        "uncertain": _pct((cards.get("uncertain") or {}).get("value", 0), rev),
    }
    return out


def dynamics(project_path: str, scope: dict | None,
             limit_days: int = 90) -> list:
    """По дням: [{day, reviewed, bad}]. Агрегация SQL, без вытягивания."""
    base = scope_filters(scope)
    codes = reviewed_codes(project_path)
    bad_codes = _codes_by_base(project_path).get("bad", [])
    if not codes:
        return []
    conds = ["COALESCE(c.hidden, 0) = 0",
             f"COALESCE(a.status, 'unreviewed') IN ({','.join(['?'] * len(codes))})"]
    params: list = list(codes)
    if base.get("file_id"):
        conds.append("c.file_id = ?")
        params.append(base["file_id"])
    if base.get("reviewed_from"):
        conds.append("substr(a.updated_at, 1, 10) >= ?")
        params.append(base["reviewed_from"])
    if base.get("reviewed_to"):
        conds.append("substr(a.updated_at, 1, 10) <= ?")
        params.append(base["reviewed_to"])
    where = "WHERE " + " AND ".join(conds)
    bph = ",".join(["?"] * len(bad_codes)) if bad_codes else ""
    q = ("SELECT substr(a.updated_at, 1, 10) AS day, "
         "COUNT(DISTINCT c.case_id) AS reviewed, "
         "COUNT(DISTINCT CASE WHEN COALESCE(a.status,'unreviewed') "
         + (f"IN ({bph}) " if bph else "= '##never##' ")
         + "THEN c.case_id END) AS bad "
         "FROM cases c "
         "JOIN files f ON c.file_id = f.file_id "
         "JOIN annotations a ON a.case_id = c.case_id "
         f"{where} "
         "GROUP BY day ORDER BY day DESC LIMIT ?")
    # ВАЖНО: порядок args — строго по порядку плейсхолдеров в строке:
    # сначала bad-IN из SELECT, потом условия WHERE, потом LIMIT.
    args = list(bad_codes) + list(params) + [int(limit_days)]
    try:
        with db(project_path) as conn:
            rows = conn.execute(q, args).fetchall()
    except Exception as e:
        logger.warning("dynamics failed: %s", e)
        return []
    return [{"day": r["day"], "reviewed": r["reviewed"], "bad": r["bad"]}
            for r in rows]


def top_categories(project_path: str, scope: dict | None,
                   limit: int = 10) -> list:
    """Частые категории ошибок (только существующие). С drill-фильтрами."""
    base = scope_filters(scope)
    conds, params = ["COALESCE(c.hidden, 0) = 0"], []
    if base.get("file_id"):
        conds.append("c.file_id = ?")
        params.append(base["file_id"])
    if base.get("reviewed_from"):
        conds.append("substr(e.updated_at, 1, 10) >= ?")
        params.append(base["reviewed_from"])
    if base.get("reviewed_to"):
        conds.append("substr(e.updated_at, 1, 10) <= ?")
        params.append(base["reviewed_to"])
    where = "WHERE " + " AND ".join(conds)
    q = (f"SELECT COALESCE(e.category_id, e.subcategory_id) AS cat, "
         "COUNT(DISTINCT e.case_id) AS n "
         "FROM case_errors e "
         "JOIN cases c ON c.case_id = e.case_id "
         "JOIN files f ON f.file_id = c.file_id "
         f"{where} GROUP BY cat ORDER BY n DESC LIMIT ?")
    try:
        with db(project_path) as conn:
            cur = conn.cursor()
            rows = cur.execute(q, params + [int(limit)]).fetchall()
            names = {r["category_id"]: r["name"] for r in cur.execute(
                "SELECT category_id, name FROM error_categories").fetchall()}
    except Exception as e:
        logger.warning("top_categories failed: %s", e)
        return []
    total = sum(r["n"] for r in rows) or 0
    out = []
    for r in rows:
        if not r["cat"]:
            continue
        flt = dict(base, error_category_id=r["cat"],
                   statuses=reviewed_codes(project_path))
        out.append({"category_id": r["cat"],
                    "name": names.get(r["cat"], f"#{r['cat']}"),
                    "problems": r["n"],
                    "share": _pct(r["n"], total),
                    "filters": flt})
    return out


def severity_dist(project_path: str, scope: dict | None) -> list:
    """Распределение тяжести (только встречающиеся; пусто — блок прячем)."""
    base = scope_filters(scope)
    conds, params = ["COALESCE(c.hidden, 0) = 0"], []
    if base.get("file_id"):
        conds.append("c.file_id = ?")
        params.append(base["file_id"])
    if base.get("reviewed_from"):
        conds.append("substr(e.updated_at, 1, 10) >= ?")
        params.append(base["reviewed_from"])
    if base.get("reviewed_to"):
        conds.append("substr(e.updated_at, 1, 10) <= ?")
        params.append(base["reviewed_to"])
    where = "WHERE " + " AND ".join(conds)
    try:
        with db(project_path) as conn:
            rows = conn.execute(
                f"SELECT e.severity AS sev, COUNT(DISTINCT e.case_id) AS n "
                f"FROM case_errors e "
                f"JOIN cases c ON c.case_id = e.case_id "
                f"JOIN files f ON f.file_id = c.file_id "
                f"{where} GROUP BY sev ORDER BY n DESC").fetchall()
    except Exception as e:
        logger.warning("severity_dist failed: %s", e)
        return []
    total = sum(r["n"] for r in rows) or 0
    names = {"low": "Низкая", "medium": "Средняя", "high": "Высокая",
             "critical": "Критическая"}
    out = []
    for r in rows:
        if not r["sev"]:
            continue
        flt = dict(base, error_severities=[r["sev"]],
                   statuses=reviewed_codes(project_path))
        out.append({"severity": r["sev"],
                    "name": names.get(r["sev"], r["sev"]),
                    "problems": r["n"],
                    "share": _pct(r["n"], total),
                    "filters": flt})
    return out


def verdict_dist(project_path: str, scope: dict | None) -> list:
    """Распределение вердиктов по base (свои коды складываются)."""
    counts = (card_counts(project_path, scope) or {}).get("cards", {})
    out = []
    for b in list(BASES) + ["unreviewed"]:
        c = counts.get(b) or {}
        out.append({"base": b, "value": c.get("value", 0) or 0,
                    "filters": c.get("filters", {})})
    return out


def bugs_stats(project_path: str, scope: dict | None) -> dict:
    """Баги за период: создано/открыто/закрыто/по тяжести/связанных кейсов."""
    base = scope_filters(scope)
    conds, params = [], []
    if base.get("reviewed_from"):
        conds.append("substr(created_at, 1, 10) >= ?")
        params.append(base["reviewed_from"])
    if base.get("reviewed_to"):
        conds.append("substr(created_at, 1, 10) <= ?")
        params.append(base["reviewed_to"])
    where = ("WHERE " + " AND ".join(conds)) if conds else ""
    try:
        with db(project_path) as conn:
            cur = conn.cursor()
            tables = {r["name"] for r in cur.execute(
                "SELECT name FROM sqlite_master WHERE type='table'").fetchall()}
            if "bug_reports" not in tables:
                return {"present": False}
            rows = cur.execute(
                f"SELECT bug_id, status, severity FROM bug_reports "
                f"{where}").fetchall()
            sev_rows = cur.execute(
                f"SELECT severity, COUNT(*) AS n FROM bug_reports "
                f"{where} GROUP BY severity").fetchall() if rows else []
            linked = 0
            if "bug_report_cases" in tables and rows:
                ph = ",".join(["?"] * len(rows))
                linked = cur.execute(
                    "SELECT COUNT(DISTINCT case_id) AS c FROM bug_report_cases "
                    f"WHERE bug_id IN ({ph})",
                    [r["bug_id"] for r in rows]).fetchone()["c"]
    except Exception as e:
        logger.warning("bugs_stats failed: %s", e)
        return {"present": False, "error": True}
    open_st = {"New", "Confirmed", "In Progress"}
    created = len(rows)
    opened = sum(1 for r in rows if (r["status"] or "") in open_st)
    out = {"present": True, "created": created, "open": opened,
           "closed": created - opened, "linked_cases": linked or 0,
           "by_severity": [{"severity": r["severity"] or "—", "n": r["n"]}
                            for r in sev_rows]}
    return out


def export_csv(project_path: str, scope: dict | None, path: str) -> str:
    """CSV аналитики с шапкой (период/фильтры/дата/проект)."""
    import csv as _csv
    from datetime import datetime as _dt
    from pathlib import Path as _P
    cards = (card_counts(project_path, scope) or {}).get("cards", {})
    pct = percentages({"cards": cards})
    dyn = dynamics(project_path, scope)
    top = top_categories(project_path, scope)
    out = _P(path)
    if out.suffix.lower() != ".csv":
        raise ValueError("Нужен .csv")
    if not out.parent.exists():
        raise FileNotFoundError(f"Папка не найдена: {out.parent}")
    with open(out, "w", encoding="utf-8-sig", newline="") as fh:
        w = _csv.writer(fh, delimiter=";")
        w.writerow(["# LocalReviewer: аналитика",
                    _P(project_path).name,
                    _dt.now().strftime("%Y-%m-%d %H:%M")])
        w.writerow(["# период",
                    (scope or {}).get("reviewed_from", "") or "—",
                    (scope or {}).get("reviewed_to", "") or "—",
                    f"файл: {(scope or {}).get('file_id', '') or 'все'}"])
        w.writerow([])
        w.writerow(["показатель", "значение"])
        for k in ("total", "reviewed", "good", "bad", "uncertain", "skip",
                  "duplicate", "unreviewed"):
            w.writerow([k, (cards.get(k) or {}).get("value", 0)])
        for k in ("good", "bad", "uncertain"):
            v = pct.get(k)
            w.writerow([f"{k} %", "" if v is None else v])
        w.writerow([])
        w.writerow(["день", "проверено", "плохих"])
        for d in dyn:
            w.writerow([d["day"], d["reviewed"], d["bad"]])
        w.writerow([])
        w.writerow(["категория", "проблем", "доля %"])
        for t in top:
            w.writerow([t["name"], t["problems"],
                        "" if t["share"] is None else t["share"]])
    logger.info("analytics csv: %s", out)
    return str(out)


def export_xlsx(project_path: str, scope: dict | None, path: str) -> str:
    """XLSX аналитики: те же разделы + шапка (период/фильтры/дата/проект)."""
    from datetime import datetime as _dt
    from pathlib import Path as _P
    from openpyxl import Workbook as _WB
    from openpyxl.styles import Font as _Font
    from export_service import _atomic_save, _resolve_output
    cards = (card_counts(project_path, scope) or {}).get("cards", {})
    pct = percentages({"cards": cards})
    dyn = dynamics(project_path, scope)
    top = top_categories(project_path, scope)
    sev = severity_dist(project_path, scope)
    out = _resolve_output(path)
    wb = _WB()
    ws = wb.active
    ws.title = "Сводка"
    bold = _Font(bold=True)
    ws.append(["LocalReviewer: аналитика", _P(project_path).name,
               _dt.now().strftime("%Y-%m-%d %H:%M")])
    ws.append(["Период:", (scope or {}).get("reviewed_from", "") or "—",
               (scope or {}).get("reviewed_to", "") or "—",
               f"файл: {(scope or {}).get('file_id', '') or 'все'}"])
    ws.append([])
    ws.append(["Показатель", "Значение"])
    for row in ws.iter_rows(min_row=4, max_row=4):
        for c in row:
            c.font = bold
    for k in ("total", "reviewed", "good", "bad", "uncertain", "skip",
              "duplicate", "unreviewed"):
        ws.append([k, (cards.get(k) or {}).get("value", 0)])
    for k in ("good", "bad", "uncertain"):
        v = pct.get(k)
        ws.append([f"{k} %", "" if v is None else v])
    ws2 = wb.create_sheet("Динамика")
    ws2.append(["День", "Проверено", "Плохих"])
    for d in dyn:
        ws2.append([d["day"], d["reviewed"], d["bad"]])
    ws3 = wb.create_sheet("Проблемы")
    ws3.append(["Категория", "Проблем", "Доля %"])
    for t in top:
        ws3.append([t["name"], t["problems"],
                    "" if t["share"] is None else t["share"]])
    ws4 = wb.create_sheet("Тяжесть")
    ws4.append(["Тяжесть", "Проблем", "Доля %"])
    for s in sev:
        ws4.append([s["name"], s["problems"],
                    "" if s["share"] is None else s["share"]])
    _atomic_save(wb, out)
    logger.info("analytics xlsx: %s", out)
    return str(out)


def error_cases_count(project_path: str, scope: dict | None) -> dict:
    """Кейсов с зафиксированной ошибкой + drill-фильтр (число == drill)."""
    from filter_service import count_filtered_cases
    base = scope_filters(scope)
    rev = reviewed_codes(project_path)
    flt = dict(base, has_errors=True, statuses=list(rev))
    return {"value": count_filtered_cases(project_path, flt), "filters": flt}


def bug_linked_case_ids(project_path: str, scope: dict | None) -> list:
    """Кейсы багов, созданных в период (для drill-down «Баги → кейсы»)."""
    base = scope_filters(scope)
    conds, params = [], []
    if base.get("reviewed_from"):
        conds.append("substr(b.created_at, 1, 10) >= ?")
        params.append(base["reviewed_from"])
    if base.get("reviewed_to"):
        conds.append("substr(b.created_at, 1, 10) <= ?")
        params.append(base["reviewed_to"])
    where = ("WHERE " + " AND ".join(conds)) if conds else ""
    try:
        with db(project_path) as conn:
            cur = conn.cursor()
            tables = {r["name"] for r in cur.execute(
                "SELECT name FROM sqlite_master WHERE type='table'").fetchall()}
            if "bug_reports" not in tables or "bug_report_cases" not in tables:
                return []
            rows = cur.execute(
                f"SELECT DISTINCT bc.case_id AS cid FROM bug_reports b "
                f"JOIN bug_report_cases bc ON bc.bug_id = b.bug_id {where}",
                params).fetchall()
    except Exception as e:
        logger.warning("bug_linked_case_ids failed: %s", e)
        return []
    return [r["cid"] for r in rows if r["cid"]]


def _scope_cards(project_path: str, scope: dict | None) -> tuple:
    """(cards, pct, errors, bugs) одним проходом для сравнения."""
    cards = (card_counts(project_path, scope) or {}).get("cards", {})
    return (cards, percentages({"cards": cards}),
            error_cases_count(project_path, scope).get("value", 0),
            bugs_stats(project_path, scope))


def _delta(a, b):
    """Разница п.п./шт: None, если считать не из чего (нет данных)."""
    if a is None or b is None:
        return None
    try:
        return round(float(b) - float(a), 1)
    except Exception:
        return None


def compare_periods(project_path: str, scope_a: dict | None,
                    scope_b: dict | None) -> dict:
    """Сравнение двух периодов: только факты, без выводов (ТЗ §9).

    Возвращает rows [{metric, a, b, delta}] + top_a/top_b (топ категорий
    каждого периода) + scopes. Дельта % — в процентных пунктах.
    """
    ca, pa, ea, ba = _scope_cards(project_path, scope_a)
    cb, pb, eb, bb = _scope_cards(project_path, scope_b)

    def _gv(cards, k):
        return (cards.get(k) or {}).get("value", 0) or 0

    na = ba.get("created") if ba.get("present") else None
    nb = bb.get("created") if bb.get("present") else None
    rows = [
        {"metric": "Проверено", "a": _gv(ca, "reviewed"),
         "b": _gv(cb, "reviewed"),
         "delta": _delta(_gv(ca, "reviewed"), _gv(cb, "reviewed"))},
        {"metric": "Good %", "a": pa.get("good"), "b": pb.get("good"),
         "delta": _delta(pa.get("good"), pb.get("good"))},
        {"metric": "Bad %", "a": pa.get("bad"), "b": pb.get("bad"),
         "delta": _delta(pa.get("bad"), pb.get("bad"))},
        {"metric": "С ошибками", "a": ea, "b": eb,
         "delta": _delta(ea, eb)},
        {"metric": "Баги", "a": na, "b": nb, "delta": _delta(na, nb)},
    ]
    return {"rows": rows,
            "top_a": top_categories(project_path, scope_a, limit=5),
            "top_b": top_categories(project_path, scope_b, limit=5),
            "scope_a": scope_a or {}, "scope_b": scope_b or {}}
