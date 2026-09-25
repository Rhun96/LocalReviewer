"""Аналитика: карточки, период, динамика, drill-down равенство, экспорт."""
import tempfile

import analytics_service as an
from database import init_database, db
from importer import import_file
from bulk_operation_service import bulk_set_status
from filter_service import get_filtered_case_ids, count_filtered_cases


def _proj():
    p = tempfile.mkdtemp()
    init_database(p)
    import_file(p, "f.xlsx", "excel", "S", 0,
                {"q": "primary_text", "i": "source_id"},
                [{"q": "q1", "i": "k1"}, {"q": "q2", "i": "k2"},
                 {"q": "q3", "i": "k3"}, {"q": "q4", "i": "k4"}])
    with db(p) as conn:
        ids = [r["case_id"] for r in conn.execute(
            "SELECT case_id FROM cases ORDER BY case_id").fetchall()]
    bulk_set_status(p, ids[:2], "good")
    bulk_set_status(p, ids[2:3], "bad")
    with db(p) as conn:
        conn.execute("UPDATE annotations SET updated_at='2020-01-01T00:00:00+00:00'"
                     " WHERE case_id=?", (ids[0],))
        conn.execute("INSERT INTO case_errors (case_id, severity, updated_at)"
                     " VALUES (?, 'high', ?)",
                     (ids[2], "2026-09-23T10:00:00+00:00"))
    return p, ids


def test_cards_and_key_drill_equality():
    """Ключевой тест ТЗ: число на карточке == кейсов после клика."""
    p, ids = _proj()
    res = an.card_counts(p, {})
    cards = res["cards"]
    assert cards["total"]["value"] == 4
    assert cards["reviewed"]["value"] == 3
    assert cards["good"]["value"] == 2
    assert cards["bad"]["value"] == 1
    assert cards["unreviewed"]["value"] == 1
    for key, card in cards.items():
        n = count_filtered_cases(p, card["filters"])
        assert n == card["value"], (key, n, card["value"])
        assert len(get_filtered_case_ids(p, card["filters"])) == card["value"]
    pct = an.percentages(res)
    assert pct["good"] == round(100 * 2 / 3, 1)
    assert pct["bad"] == round(100 * 1 / 3, 1)


def test_period_filters_and_empty_states():
    p, ids = _proj()
    res = an.card_counts(p, {"reviewed_from": "2026-01-01"})
    assert res["cards"]["reviewed"]["value"] == 2  # без позапрошлогоднего
    assert res["cards"]["good"]["value"] == 1
    res = an.card_counts(p, {"reviewed_from": "2030-01-01"})
    assert res["cards"]["reviewed"]["value"] == 0
    assert an.percentages(res) == {"good": None, "bad": None,
                                   "uncertain": None}
    assert an.dynamics(p, {"reviewed_from": "2030-01-01"}) == []
    # drill по периоду тоже сходится
    flt = res["cards"]["bad"]["filters"]
    assert flt["reviewed_from"] == "2030-01-01"
    assert count_filtered_cases(p, flt) == 0


def test_dynamics_groups_by_day():
    p, ids = _proj()
    dyn = an.dynamics(p, {})
    assert len(dyn) >= 2
    total = sum(d["reviewed"] for d in dyn)
    assert total == 3
    assert sum(d["bad"] for d in dyn) == 1
    days = [d["day"] for d in dyn]
    assert days == sorted(days, reverse=True)
    assert "2020-01-01" in days


def test_top_categories_and_severity():
    p, ids = _proj()
    top = an.top_categories(p, {})
    assert top == []  # категории не заданы — честно пусто, не нули
    sev = an.severity_dist(p, {})
    assert len(sev) == 1 and sev[0]["severity"] == "high"
    assert sev[0]["problems"] == 1
    assert count_filtered_cases(p, sev[0]["filters"]) == 1
    # период отсекает ошибку
    assert an.severity_dist(p, {"reviewed_from": "2030-01-01"}) == []


def test_verdict_dist_and_bugs():
    p, ids = _proj()
    dist = {d["base"]: d["value"] for d in an.verdict_dist(p, {})}
    assert dist == {"good": 2, "bad": 1, "uncertain": 0, "skip": 0,
                    "duplicate": 0, "unreviewed": 1}
    bugs = an.bugs_stats(p, {})
    assert bugs["present"] is True and bugs["created"] == 0
    import bug_report_service as bugs_svc
    bugs_svc.create_bug(p, "t", [ids[2]], severity="High")
    bugs = an.bugs_stats(p, {})
    assert (bugs["created"], bugs["open"], bugs["linked_cases"]) == (1, 1, 1)
    assert bugs["by_severity"] == [{"severity": "High", "n": 1}]


def test_export_csv_with_header():
    import csv as _csv
    p, ids = _proj()
    out = p + "/analytics.csv"
    got = an.export_csv(p, {"reviewed_from": "2026-01-01"}, out)
    import os as _os
    assert _os.path.normcase(got) == _os.path.normcase(out)
    with open(out, encoding="utf-8-sig") as fh:
        rows = list(_csv.reader(fh, delimiter=";"))
    assert rows[0][0].startswith("# LocalReviewer")
    assert "2026-01-01" in rows[1]
    flat = [c for r in rows for c in r]
    assert "bad" in flat and "день" in flat


def test_export_xlsx_sections():
    import tempfile
    from openpyxl import load_workbook
    p, ids = _proj()
    out = tempfile.mkdtemp() + "/analytics.xlsx"
    got = an.export_xlsx(p, {}, out)
    import os as _os
    assert _os.path.normcase(got) == _os.path.normcase(out)
    wb = load_workbook(out, read_only=True, data_only=True)
    assert wb.sheetnames == ["Сводка", "Динамика", "Проблемы", "Тяжесть"]
    ws = wb["Сводка"]
    assert ws.cell(row=1, column=1).value.startswith("LocalReviewer")
    assert ws.cell(row=2, column=1).value == "Период:"


def test_case_ids_drill_filter():
    p, ids = _proj()
    flt = {"case_ids": ids[:2]}
    assert count_filtered_cases(p, flt) == 2
    assert get_filtered_case_ids(p, flt) == ids[:2]
    big = {"case_ids": list(range(1, 1200))}
    assert count_filtered_cases(p, big) == 4  # чанки IN, лишних нет


def test_compare_periods_facts_only():
    """A vs B: факты и дельты, без выводов; пустой период — прочерки."""
    p, ids = _proj()
    res = an.compare_periods(p, {"reviewed_from": "2020-01-01",
                                 "reviewed_to": "2020-12-31"},
                             {"reviewed_from": "2026-01-01"})
    by_m = {r["metric"]: r for r in res["rows"]}
    assert by_m["Проверено"] == {"metric": "Проверено", "a": 1, "b": 2,
                                 "delta": 1.0}
    assert by_m["Bad %"]["a"] == 0.0
    assert by_m["Bad %"]["b"] == 50.0
    assert by_m["Bad %"]["delta"] == 50.0
    assert len(res["top_a"]) == 0  # в 2020 ошибок не было — честно пусто
    empty = an.compare_periods(p, {"reviewed_from": "2030-01-01"}, {})
    em = {r["metric"]: r for r in empty["rows"]}
    assert em["Good %"] == {"metric": "Good %", "a": None, "b": 66.7,
                            "delta": None}


def test_file_scope_limits_summary():
    """Охват файлом режет и карточки, и динамику, и топ."""
    from database import db
    p = tempfile.mkdtemp()
    init_database(p)
    import_file(p, "a.xlsx", "excel", "S", 0,
                {"q": "primary_text", "i": "source_id"},
                [{"q": "a1", "i": "a1"}, {"q": "a2", "i": "a2"}])
    import_file(p, "b.xlsx", "excel", "S", 0,
                {"q": "primary_text", "i": "source_id"},
                [{"q": "b1", "i": "b1"}])
    with db(p) as conn:
        fids = {r["file_name"]: r["file_id"] for r in conn.execute(
            "SELECT file_id, file_name FROM files").fetchall()}
        aids = [r["case_id"] for r in conn.execute(
            "SELECT case_id FROM cases WHERE file_id=?", (fids["a.xlsx"],)).fetchall()]
    bulk_set_status(p, aids, "good")
    all_cards = an.card_counts(p, {})["cards"]
    assert all_cards["reviewed"]["value"] == 2
    scoped = an.card_counts(p, {"file_id": fids["b.xlsx"]})["cards"]
    assert scoped["reviewed"]["value"] == 0
    assert scoped["total"]["value"] == 1
    dyn = an.dynamics(p, {"file_id": fids["a.xlsx"]})
    assert sum(d["reviewed"] for d in dyn) == 2
    assert an.top_categories(p, {"file_id": fids["b.xlsx"]}) == []
