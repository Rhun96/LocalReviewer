"""Этап B: прогоны модели — импорт, сопоставление, сравнение, предпочтения."""
import tempfile

import pytest
from database import init_database
from importer import import_file
import model_run_service as m


def _proj():
    tmp = tempfile.mkdtemp()
    init_database(tmp)
    import_file(tmp, "base.xlsx", "excel", "S", 0,
                {"q": "primary_text", "i": "source_id"},
                [{"q": "можно вернуть билет?", "i": "k1"},
                 {"q": "где мой заказ", "i": "k2"}])
    return tmp


def test_run_crud_and_validation():
    p = _proj()
    a = m.create_run(p, "v1.7", "model-x", "1.7", "p3")
    assert a > 0
    with pytest.raises(ValueError):
        m.create_run(p, "v1.7", "model-x")  # дубль имени
    with pytest.raises(ValueError):
        m.create_run(p, "vX", "")  # без модели
    assert len(m.list_runs(p)) == 1
    m.delete_run(p, a)
    assert m.list_runs(p) == []
    with pytest.raises(ValueError):
        m.delete_run(p, a)


def test_import_matching_and_dups():
    p = _proj()
    a = m.create_run(p, "v1", "mx")
    res = m.import_run_rows(p, a, [
        {"source_id": "k1", "answer": "да", "prompt": "можно вернуть билет?"},
        {"source_id": "k1", "answer": "да-да", "prompt": "можно вернуть билет?"},
        {"source_id": "", "answer": "там", "prompt": "где мой заказ"},
        {"source_id": "", "answer": "???", "prompt": ""},
        {"source_id": "k9", "answer": "новое", "prompt": "новый вопрос"},
    ])
    assert res["total"] == 5
    assert res["matched"] == 2  # k1 + текст k2
    assert res["new"] == 1  # k9
    assert res["dups"] == 1  # повтор k1
    assert res["no_key"] == 1  # без ID и промпта
    answers = {x["stable_key"]: x for x in m.list_answers(p, a)}
    assert answers["src:k1"]["case_id"] is not None
    assert answers["src:k1"]["answer_text"] == "да"  # первый выигрывает
    assert answers["src:k9"]["case_id"] is None


def test_compare_and_preferences():
    p = _proj()
    a = m.create_run(p, "A", "mx")
    b = m.create_run(p, "B", "mx")
    m.import_run_rows(p, a, [
        {"source_id": "k1", "answer": "ответ A1"},
        {"source_id": "k2", "answer": "ответ A2"}])
    m.import_run_rows(p, b, [
        {"source_id": "k1", "answer": "ответ B1"},
        {"source_id": "k3", "answer": "ответ B3"}])
    data = m.compare_runs(p, a, b)
    assert data["counts"]["both"] == 1
    assert data["counts"]["only_a"] == 1
    assert data["counts"]["only_b"] == 1
    both = next(r for r in data["rows"] if r["stable_key"] == "src:k1")
    assert both["answer_a"] == "ответ A1" and both["answer_b"] == "ответ B1"
    assert both["case_status"] == "unreviewed"  # absolute из разметки
    assert both["verdict"] == "unknown" and not both["has_verdict"]
    m.set_preference(p, a, b, "src:k1", "b_better", 2, 1, "точнее")
    data2 = m.compare_runs(p, a, b)
    both2 = next(r for r in data2["rows"] if r["stable_key"] == "src:k1")
    assert both2["verdict"] == "b_better" and both2["rank_a"] == 2
    assert both2["comment"] == "точнее"
    assert m.preference_stats(p, a, b)["b_better"] == 1
    with pytest.raises(ValueError):
        m.set_preference(p, a, b, "src:k1", "nope")
    with pytest.raises(ValueError):
        m.set_preference(p, a, b, "src:k1", "tie", 5, None)
    with pytest.raises(ValueError):
        m.compare_runs(p, a, 9999)


def test_roundrobin_matrix_and_standings():
    """Круговая таблица: все пары разом, места из предпочтений."""
    p = _proj()
    a = m.create_run(p, "A", "mx")
    b = m.create_run(p, "B", "mx")
    c = m.create_run(p, "C", "mx")
    for r in (a, b, c):
        m.import_run_rows(p, r, [
            {"source_id": "k1", "answer": f"ans {r} 1"},
            {"source_id": "k2", "answer": f"ans {r} 2"}])
    m.set_preference(p, a, b, "src:k1", "a_better")
    m.set_preference(p, a, b, "src:k2", "tie")
    m.set_preference(p, b, c, "src:k1", "b_better")
    m.set_preference(p, c, a, "src:k2", "a_better")
    res = m.roundrobin_matrix(p, [a, b, c])
    assert [r["run_id"] for r in res["runs"]] == [a, b, c]
    assert len(res["pairs"]) == 3
    by_pair = {(x["a"], x["b"]): x for x in res["pairs"]}
    ab = by_pair[(a, b)]
    assert (ab["wins_a"], ab["wins_b"], ab["ties"]) == (1, 0, 1)
    assert ab["common"] == 2 and ab["unjudged"] == 0
    bc = by_pair[(b, c)]
    assert bc["wins_b"] == 1 and bc["unjudged"] == 1
    assert [s["run_id"] for s in res["standings"]] == [c, a, b]
    assert res["standings"][0]["score"] == 2.0
    with pytest.raises(ValueError):
        m.roundrobin_matrix(p, [a])
    with pytest.raises(ValueError):
        m.roundrobin_matrix(p, [a, 9999])


def test_run_answer_product_and_metadata():
    import json as _json
    p = _proj()
    a = m.create_run(p, "v", "mx")
    res = m.import_run_rows(p, a, [
        {"source_id": "k1", "answer": "да", "product": "Авиа",
         "metadata": {"Канал": "Чат", "пусто": "  "}},
        {"source_id": "k2", "answer": "там"},
    ])
    assert res["matched"] == 2
    ans = {x["stable_key"]: x for x in m.list_answers(p, a)}
    assert ans["src:k1"]["product"] == "Авиа"
    assert _json.loads(ans["src:k1"]["metadata_json"]) == {"Канал": "Чат"}
    assert ans["src:k2"]["product"] in ("", None)
    assert ans["src:k2"]["metadata_json"] is None
    b = m.create_run(p, "w", "mx")
    m.import_run_rows(p, b, [{"source_id": "k1", "answer": "да2"}])
    row = next(r for r in m.compare_runs(p, a, b)["rows"]
               if r["stable_key"] == "src:k1")
    assert row["product_a"] == "Авиа" and row["product_b"] == ""
    assert row["product_case"] == ""


def test_run_custom_metadata_grouping():
    from runs_screen import collect_run_metadata
    row = {"a": "x", "b": "y", "c": " ", "d": "z"}
    assert collect_run_metadata(row, {"a": "K", "b": "K", "c": "K"}) == {"K": "x\n\ny"}
    assert collect_run_metadata(row, {"d": "D2"}) == {"D2": "z"}
    assert collect_run_metadata(row, {}) == {}


def test_run_rows_without_answer_column():
    p = _proj()
    a = m.create_run(p, "q-only", "mx")
    res = m.import_run_rows(p, a, [
        {"source_id": "k1"},
        {"source_id": "k2", "answer": ""},
        {"source_id": "", "prompt": "где мой заказ"},
    ])
    assert res["matched"] == 3 and res["no_key"] == 0
    ans = {x["stable_key"]: x for x in m.list_answers(p, a)}
    assert ans["src:k1"]["answer_text"] is None


def test_runs_listed_oldest_first():
    p = _proj()
    a = m.create_run(p, "first", "mx")
    b = m.create_run(p, "second", "mx")
    assert [r["run_id"] for r in m.list_runs(p)] == [a, b]


def test_run_answer_prompt_text_fallback():
    p = _proj()
    a = m.create_run(p, "v", "mx")
    m.import_run_rows(p, a, [{"answer": "x", "prompt": "новый вопрос без кейса"}])
    ans = m.list_answers(p, a)
    assert len(ans) == 1 and ans[0]["case_id"] is None
    assert ans[0]["prompt_text"] == "новый вопрос без кейса"


def test_import_cancel_rolls_back():
    import threading
    from workers import Cancelled
    p = _proj()
    a = m.create_run(p, "v", "mx")
    ev = threading.Event()
    ev.set()
    with pytest.raises(Cancelled):
        m.import_run_rows(p, a, [{"source_id": "k1", "answer": "x"}],
                          cancel_event=ev)
    assert m.list_answers(p, a) == []


def test_compare_side_marks_and_severity():
    """Сравнение отдаёт статусы/комменты сторон и тяжесть кейса."""
    import regression_service as rg
    p = _proj()
    a = m.create_run(p, "A", "mx")
    b = m.create_run(p, "B", "mx")
    m.import_run_rows(p, a, [{"source_id": "k1", "answer": "a1"}])
    m.import_run_rows(p, b, [{"source_id": "k1", "answer": "b1"}])
    rg.set_output_review(p, a, "src:k1", "bad", "Критичность: высокая\nплохо")
    rg.set_output_review(p, b, "src:k1", "good", "норм")
    data = m.compare_runs(p, a, b)
    row = next(r for r in data["rows"] if r["stable_key"] == "src:k1")
    assert row["status_a"] == "bad" and row["status_b"] == "good"
    assert "высокая" in (row["comment_a"] or "")
    from compare_dialog import _side_severity, _side_is_high, _side_is_bad
    assert _side_severity(row["comment_a"], "") == "высокая"
    assert _side_severity("", "high") == "высокая"
    assert _side_is_high("высокая") and _side_is_high("high")
    assert not _side_is_high("низкая")
    assert _side_is_bad(p, "bad") and not _side_is_bad(p, "good")
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from compare_dialog import CompareDialog
    dlg = CompareDialog(p, a, b, None)
    try:
        dlg.show()
        assert "Критичность: высокая" in dlg.pane_a.toPlainText()
        assert "Плохих A: 1" in dlg.stats.text(), dlg.stats.text()
        assert "Δ Bad" in dlg.stats.text(), dlg.stats.text()
    finally:
        dlg.close()


def test_bulk_review_marks_cells_and_copies():
    """Массовая разметка: ячейка, автопереход, копия на все ответы кейса."""
    import os
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    import regression_service as rg
    from bulk_review_dialog import BulkReviewDialog
    p = _proj()
    a = m.create_run(p, "A", "mx")
    b = m.create_run(p, "B", "mx")
    for r in (a, b):
        m.import_run_rows(p, r, [{"source_id": "k1", "answer": "xa"},
                                 {"source_id": "k2", "answer": "xb"}])
    d = BulkReviewDialog(p, [a, b], None)
    try:
        d.show()
        assert len(d._keys) == 2 and len(d.panes) == 2
        d.keys_list.setCurrentRow(0)
        d._set_status("good")
        assert rg.output_review_stats(p, a)["reviewed"] == 1
        assert d._active == 1 and d._current_key() == "src:k1"
        d._select_run(a)
        d._apply_to_all()
        assert rg.output_review_stats(p, b)["reviewed"] == 1
        assert "2/4" in d.title.text(), d.title.text()
    finally:
        d.close()


def test_bulk_review_comment_autosaves_like_review():
    """Коммент живёт без статуса: уход с ячейки сохраняет, как в ревью."""
    import os
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    import regression_service as rg
    from bulk_review_dialog import BulkReviewDialog
    p = _proj()
    a = m.create_run(p, "A", "mx")
    m.import_run_rows(p, a, [{"source_id": "k1", "answer": "xa"},
                             {"source_id": "k2", "answer": "xb"}])
    rg.set_output_review(p, a, "src:k1", "good", "старый коммент")
    d = BulkReviewDialog(p, [a], None)
    try:
        d.show()
        d.keys_list.setCurrentRow(0)
        # открытие не затёрло коммент пустотой
        assert d.comment_edit.toPlainText() == "старый коммент"
        assert rg.list_output_reviews(p, a)[0]["review_comment"] == \
            "старый коммент"
        # набрали новый, ушли без статуса — сохранился, статус цел
        d.comment_edit.setPlainText("новый коммент")
        d.keys_list.setCurrentRow(1)
        rows = {r["stable_key"]: r
                for r in rg.list_output_reviews(p, a)}
        assert rows["src:k1"]["review_comment"] == "новый коммент"
        assert rows["src:k1"]["review_status"] == "good"
        assert rows["src:k2"]["review_comment"] in (None, "")
    finally:
        d.close()


def test_run_surfaces_show_case_topic():
    """Тема видна при разметке/сравнении прогонов (и своя «Тема»)."""
    import os
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from importer import import_file
    import tempfile
    from database import init_database
    p = tempfile.mkdtemp()
    init_database(p)
    import_file(p, "t.xlsx", "excel", "S", 0,
                {"q": "primary_text", "t": "custom:Тема", "i": "source_id"},
                [{"q": "где билет", "t": "Возвраты", "i": "k1"}])
    a = m.create_run(p, "A", "mx")
    b = m.create_run(p, "B", "mx")
    for r in (a, b):
        m.import_run_rows(p, r, [{"source_id": "k1", "answer": "ans"}])
    assert m.compare_runs(p, a, b)["rows"][0]["topic_case"] == "Возвраты"
    from run_review_dialog import RunReviewDialog
    d1 = RunReviewDialog(p, a, None)
    try:
        d1.show()
        assert "Возвраты" in d1.prompt_label.text(), \
            d1.prompt_label.text()
    finally:
        d1.close()
    from bulk_review_dialog import BulkReviewDialog
    d2 = BulkReviewDialog(p, [a, b], None)
    try:
        d2.show()
        assert "Возвраты" in d2.prompt_label.text(), \
            d2.prompt_label.text()
    finally:
        d2.close()
    from compare_dialog import CompareDialog
    d3 = CompareDialog(p, a, b, None)
    try:
        d3.show()
        assert "Возвраты" in d3.prompt_label.text(), \
            d3.prompt_label.text()
    finally:
        d3.close()


def test_export_run_marks_roundtrip():
    """Выгрузка разметки читается импортом 1-в-1 (круг без потерь)."""
    import os
    import regression_service as rg
    import export_service as ex
    from run_marks_io_service import read_marks_table, preview_marks
    p = _proj()
    a = m.create_run(p, "A", "mx")
    m.import_run_rows(p, a, [{"source_id": "k1", "answer": "xa"},
                             {"source_id": "k2", "answer": "xb"}])
    rg.set_output_review(p, a, "src:k1", "bad", "Критичность: высокая\nплохо")
    rg.set_output_review(p, a, "src:k2", "good", None)
    path = os.path.join(p, "marks.xlsx")
    out = ex.export_run_marks(p, a, path)
    assert os.path.exists(out)
    from openpyxl import load_workbook
    wb = load_workbook(out)
    ws = wb.active
    assert [c.value for c in ws[1]] == ["ID", "Вопрос", "Ответ", "Статус",
                                        "Комментарий", "Тяжесть"]
    assert ws.max_row == 3
    headers, rows, errors = read_marks_table(out)
    assert not errors
    mapping = {"id": "ID", "status": "Статус", "comment": "Комментарий",
               "severity": "Тяжесть"}
    rep = preview_marks(p, a, rows, mapping)
    assert rep["errors"] == [] and rep["not_found"] == []
    assert len(rep["unchanged"]) == 2
    import pytest as _pt
    with _pt.raises(ValueError):
        ex.export_run_marks(p, 999999, os.path.join(p, "x.xlsx"))


def test_export_run_marks_multi_sheet():
    """Мульти-выгрузка: по листу на прогон, каждый читается импортом."""
    import os
    import export_service as ex
    from run_marks_io_service import read_marks_table, preview_marks
    p = _proj()
    a = m.create_run(p, "A", "mx")
    b = m.create_run(p, "B", "mx")
    m.import_run_rows(p, a, [{"source_id": "k1", "answer": "xa"}])
    m.import_run_rows(p, b, [{"source_id": "k1", "answer": "xb"},
                             {"source_id": "k2", "answer": "xc"}])
    import regression_service as rg
    rg.set_output_review(p, b, "src:k1", "bad", None)
    path = os.path.join(p, "multi.xlsx")
    out = ex.export_run_marks(p, [a, b], path)
    assert os.path.exists(out)
    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert wb.sheetnames == ["A", "B"]
    assert wb["A"].max_row == 2 and wb["B"].max_row == 3
    for sheet, rid in (("A", a), ("B", b)):
        headers, rows, errors = read_marks_table(out, sheet)
        assert not errors
        mapping = {"id": "ID", "status": "Статус", "comment": "Комментарий",
                   "severity": "Тяжесть"}
        rep = preview_marks(p, rid, rows, mapping)
        assert rep["errors"] == [] and rep["not_found"] == []
        n = len(rows)
        assert len(rep["unchanged"]) == n, (sheet, rep)
    import pytest as _pt
    with _pt.raises(ValueError):
        ex.export_run_marks(p, [], os.path.join(p, "e.xlsx"))
