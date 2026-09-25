"""Аналитика: продукты, топ причин, скорость/прогноз, лидерборд, отчёт руководству."""
import tempfile
from pathlib import Path

import openpyxl
import pytest

from bulk_operation_service import bulk_set_status
from database import db, init_database
from filter_service import get_filtered_case_ids
from importer import find_file_by_name, import_file


def _proj():
    tmp = tempfile.mkdtemp()
    init_database(tmp)
    m = {"q": "primary_text", "p": "product", "i": "source_id"}
    import_file(tmp, "a.xlsx", "excel", "S", 0, m,
                [{"q": "qa1", "p": "Авиа", "i": "k1"},
                 {"q": "qa2", "p": "ЖД", "i": "k2"}])
    import_file(tmp, "b.xlsx", "excel", "S", 0, m,
                [{"q": "qb1", "p": "Авиа", "i": "k3"}])
    return tmp


def _fid(proj, name):
    return find_file_by_name(proj, name)["file_id"]


def test_product_and_error_top():
    from taxonomy_service import set_case_error
    p = _proj()
    ids = get_filtered_case_ids(p, {})
    bulk_set_status(p, ids[:2], "good")
    bulk_set_status(p, ids[2:], "bad")
    with db(p) as conn:
        cat = conn.execute("SELECT category_id FROM error_categories "
                           "WHERE code='correctness'").fetchone()
        sub = conn.execute("SELECT category_id FROM error_categories "
                           "WHERE code='correctness.hallucination'").fetchone()
    set_case_error(p, ids[2], cat["category_id"], sub["category_id"], "high")
    from report_service import get_error_top, get_product_report
    prods = {r["product"]: r for r in get_product_report(p)}
    assert prods["Авиа"]["total"] == 2 and prods["Авиа"]["reviewed"] == 2
    assert prods["Авиа"]["bad"] == 1 and prods["ЖД"]["bad"] == 0
    assert [r["product"] for r in get_product_report(p, _fid(p, "a.xlsx"))]
    top = get_error_top(p)
    assert top and top[0]["category"] == "Правильность" and top[0]["n"] == 1
    assert get_error_top(p, _fid(p, "a.xlsx")) == []


def test_velocity_and_forecast():
    from report_service import get_velocity
    p = _proj()
    ids = get_filtered_case_ids(p, {})
    bulk_set_status(p, ids[:2], "good")
    with db(p) as conn:
        conn.execute("UPDATE history SET created_at='2026-01-05 10:00:00' "
                     "WHERE case_id=?", (ids[0],))
        conn.execute("UPDATE history SET created_at='2026-01-06 10:00:00' "
                     "WHERE case_id=?", (ids[1],))
    v = get_velocity(p)
    assert (v["total"], v["reviewed"], v["remaining"]) == (3, 2, 1)
    assert v["avg_per_day"] == 1.0 and v["active_days"] == 2
    assert v["eta_days"] == 1
    assert [d["day"] for d in v["per_day"]] == ["2026-01-06", "2026-01-05"]


def test_leaderboard_and_management_report():
    import model_run_service as m
    import regression_service as rg
    from export_service import export_management_report
    from report_service import get_model_leaderboard
    p = _proj()
    a = m.create_run(p, "A", "mx")
    b = m.create_run(p, "B", "mx")
    m.import_run_rows(p, a, [{"source_id": "k1", "answer": "a1"},
                             {"source_id": "k2", "answer": "a2"},
                             {"source_id": "k3", "answer": "a3"}])
    m.import_run_rows(p, b, [{"source_id": "k1", "answer": "b1"},
                             {"source_id": "k2", "answer": "b2"},
                             {"source_id": "k3", "answer": "b3"}])
    rg.set_output_review(p, a, "src:k1", "good", None)
    rg.set_output_review(p, a, "src:k2", "good", None)
    rg.set_output_review(p, a, "src:k3", "bad", None)
    rg.set_output_review(p, b, "src:k1", "bad", None)
    rg.set_output_review(p, b, "src:k2", "good", None)
    rg.set_output_review(p, b, "src:k3", "good", None)
    m.set_preference(p, a, b, "src:k1", "a_better", 1, 2, None)
    rg.run_regression(p, "r1", "run", a, "run", b)
    board = {r["run_id"]: r for r in get_model_leaderboard(p)}
    assert (board[a]["answers"], board[a]["reviewed"]) == (3, 3)
    assert (board[a]["good"], board[a]["bad"]) == (2, 1)
    assert (board[a]["wins"], board[a]["pairs"]) == (1, 1)
    assert (board[b]["gates"], board[b]["gates_passed"]) == (1, 0)
    ids = get_filtered_case_ids(p, {})
    bulk_set_status(p, ids, "good")
    out = str(Path(p) / "mgmt.xlsx")
    assert export_management_report(p, out) is True
    ws = openpyxl.load_workbook(out, read_only=True)
    assert ws.sheetnames == ["Сводка", "Сверка", "Продукты", "Причины", "Прогоны",
                             "Регрессии", "Версии"]
    summary = {r[0]: r[1] for r in ws["Сводка"].iter_rows(values_only=True)}
    assert summary["Всего кейсов"] == 3 and summary["Осталось"] == 0
    assert summary["ВЕРДИКТ"] == "НЕ ГОТОВ (gate FAIL)"


def test_consistency_alerts_golden_trend():
    from report_service import (consistency_check, get_alerts,
                                get_golden_info, get_quality_trend)
    from dataset_service import (create_dataset, create_version,
                                 freeze_version)
    p = _proj()
    chk = consistency_check(p)
    assert chk["ok"] and len(chk["checks"]) == 4
    assert get_alerts(p) == []
    with db(p) as conn:
        conn.execute("UPDATE cases SET created_at='2026-01-01 00:00:00'")
    assert any("30" in a["text"] for a in get_alerts(p)
               if a["level"] == "warning")
    ids = get_filtered_case_ids(p, {})
    bulk_set_status(p, ids[:2], "good")
    ds = create_dataset(p, "G")
    _v1 = create_version(p, ds)
    bulk_set_status(p, ids[2:], "bad")
    v2 = create_version(p, ds)
    freeze_version(p, v2)
    gi = get_golden_info(p)
    assert gi["count"] == 1 and gi["oldest_days"] in (0, 1)
    tr = {t["version"]: t for t in get_quality_trend(p)}
    assert tr[1]["bad_rate"] == 0
    assert tr[2]["bad_rate"] == pytest.approx(1 / 3, abs=1e-3)
    assert any(a["level"] == "info" and "эталон" in a["text"]
               for a in get_alerts(p)) is False
