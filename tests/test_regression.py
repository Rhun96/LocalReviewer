"""Этап C: разметка ответов, матрица, gate, фильтры, экспорт."""
import tempfile

import pytest
from database import init_database
from importer import import_file
from bulk_operation_service import bulk_set_status
from filter_service import get_filtered_case_ids
from dataset_service import create_dataset, create_version, freeze_version
import model_run_service as m
import regression_service as rg


def _proj(n=6):
    tmp = tempfile.mkdtemp()
    init_database(tmp)
    import_file(tmp, "b.xlsx", "excel", "S", 0,
                {"q": "primary_text", "i": "source_id"},
                [{"q": f"вопрос {i}", "i": f"k{i}"} for i in range(n)])
    return tmp


def _run_with_answers(p, name="cand", keys=(0, 1, 2, 3, 4, 5)):
    c = m.create_run(p, name, "mx")
    m.import_run_rows(
        p, c, [{"source_id": f"k{i}", "answer": f"ответ {i}"} for i in keys])
    return c


def test_output_review_crud_and_stats():
    p = _proj()
    c = _run_with_answers(p)
    st = rg.output_review_stats(p, c)
    assert (st["total"], st["reviewed"], st["remaining"]) == (6, 0, 6)
    rg.set_output_review(p, c, "src:k0", "good", "ок")
    rg.set_output_review(p, c, "src:k1", "bad")
    st = rg.output_review_stats(p, c)
    assert (st["reviewed"], st["remaining"]) == (2, 4)
    rows = {r["stable_key"]: r for r in rg.list_output_reviews(p, c)}
    assert rows["src:k0"]["review_status"] == "good"
    assert rows["src:k0"]["answer_text"] == "ответ 0"
    with pytest.raises(ValueError):
        rg.set_output_review(p, c, "src:nope", "good")
    with pytest.raises(ValueError):
        rg.set_output_review(p, c, "src:k0", "плохо!")


def test_matrix_and_gate():
    p = _proj()
    ids = get_filtered_case_ids(p, {})
    bulk_set_status(p, ids, "good")
    ds = create_dataset(p, "T", dataset_type="test")
    v = create_version(p, ds)
    freeze_version(p, v)
    c = _run_with_answers(p)
    rg.set_output_review(p, c, "src:k0", "bad")  # good -> bad: critical
    rg.set_output_review(p, c, "src:k1", "uncertain")  # good -> uncertain: warning
    rg.set_output_review(p, c, "src:k2", "good")
    rid = rg.run_regression(p, "r1", "dataset_version", v, "run", c,
                            gate_max_critical=0, gate_max_rate=0.5)
    g = rg.get_regression(p, rid)
    assert g["regressions"] == 2 and g["gate_result"] == "FAIL"
    rows = {r["stable_key"]: r for r in rg.list_regression_results(p, rid)}
    assert rows["src:k0"]["result"] == "REGRESSION"
    assert rows["src:k0"]["severity"] == "critical"
    assert rows["src:k1"]["severity"] == "warning"
    assert rows["src:k2"]["result"] == "UNCHANGED"
    assert rows["src:k3"]["result"] == "UNRESOLVED"  # ответ есть, разметки нет
    # мягкий gate — PASS
    rid2 = rg.run_regression(p, "r2", "dataset_version", v, "run", c,
                             gate_max_critical=5, gate_max_rate=0.5)
    assert rg.get_regression(p, rid2)["gate_result"] == "PASS"


def test_baseline_from_run_and_new_removed():
    p = _proj(4)
    a = _run_with_answers(p, "A", keys=(0, 1))
    b = _run_with_answers(p, "B", keys=(1, 2))
    for key, st in (("src:k0", "good"), ("src:k1", "bad")):
        rg.set_output_review(p, a, key, st)
    for key, st in (("src:k1", "good"), ("src:k2", "good")):
        rg.set_output_review(p, b, key, st)
    rid = rg.run_regression(p, "r", "run", a, "run", b)
    rows = {r["stable_key"]: r["result"] for r in rg.list_regression_results(p, rid)}
    assert rows == {"src:k0": "REMOVED", "src:k1": "IMPROVED", "src:k2": "NEW"}
    only = rg.list_regression_results(p, rid, result="IMPROVED")
    assert len(only) == 1 and only[0]["stable_key"] == "src:k1"
    crit = rg.list_regression_results(p, rid, severity="critical")
    assert crit == []


def test_custom_base_in_matrix():
    from review_profile_service import create_profile, set_active_profile
    p = _proj(2)
    pid = create_profile(p, "S", {
        "statuses": [
            {"code": "safe", "name": "Safe", "hotkey": "1",
             "enabled": True, "base": "good"},
            {"code": "target", "name": "Target", "hotkey": "2",
             "enabled": True, "base": "bad"}],
        "require_category_for_bad": False, "require_comment_for_bad": None})
    set_active_profile(p, pid)
    ids = get_filtered_case_ids(p, {})
    bulk_set_status(p, [ids[0]], "safe")
    bulk_set_status(p, [ids[1]], "safe")
    ds = create_dataset(p, "T", dataset_type="test")
    v = create_version(p, ds)
    c = _run_with_answers(p, keys=(0, 1))
    rg.set_output_review(p, c, "src:k0", "target")  # good -> bad
    rg.set_output_review(p, c, "src:k1", "safe")
    rid = rg.run_regression(p, "r", "dataset_version", v, "run", c)
    rows = {r["stable_key"]: r["result"] for r in rg.list_regression_results(p, rid)}
    assert rows["src:k0"] == "REGRESSION"
    assert rows["src:k1"] == "UNCHANGED"


def test_delete_and_export():
    import os
    p = _proj(2)
    ids = get_filtered_case_ids(p, {})
    bulk_set_status(p, ids, "good")
    ds = create_dataset(p, "T", dataset_type="test")
    v = create_version(p, ds)
    c = _run_with_answers(p, keys=(0, 1))
    rg.set_output_review(p, c, "src:k0", "bad")
    rg.set_output_review(p, c, "src:k1", "good")
    rid = rg.run_regression(p, "r", "dataset_version", v, "run", c)
    out = os.path.join(p, "regressions.xlsx")
    path = rg.export_regressions_xlsx(p, rid, out)
    assert os.path.exists(path)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb.active
    assert ws.max_row == 2  # заголовок + 1 регрессия
    assert ws.cell(row=1, column=1).value == "case_id"
    assert ws.cell(row=2, column=7).value.startswith("REGRESSION")
    assert [r["regression_id"] for r in rg.list_regressions(p)] == [rid]
    rg.delete_regression(p, rid)
    assert rg.list_regressions(p) == []
    with pytest.raises(ValueError):
        rg.delete_regression(p, rid)


def test_dataset_vs_dataset_candidate():
    """Кандидат — версия датасета: матрица, gate, label, экспорт."""
    p = _proj(3)
    ids = get_filtered_case_ids(p, {})
    bulk_set_status(p, ids, "good")
    ds = create_dataset(p, "T", dataset_type="test")
    v1 = create_version(p, ds)
    bulk_set_status(p, ids[:1], "bad")
    v2 = create_version(p, ds)
    rid = rg.run_regression(p, "ds-vs-ds", "dataset_version", v1,
                            "dataset_version", v2)
    g = rg.get_regression(p, rid)
    assert g["candidate_type"] == "dataset_version"
    assert g["candidate_run_id"] is None
    rows = {r["stable_key"]: r for r in rg.list_regression_results(p, rid)}
    bad_rows = [r for r in rows.values() if r["result"] == "REGRESSION"]
    assert len(bad_rows) == 1 and bad_rows[0]["severity"] == "critical"
    assert "Датасет" in rg.candidate_label(p, g)
    import os
    out = os.path.join(p, "ds_regressions.xlsx")
    assert os.path.exists(rg.export_regressions_xlsx(p, rid, out))
    # баг из строки версии-кандидата — без ответов прогона, но с контекстом
    pre = rg.bug_prefill(p, rid, bad_rows[0]["stable_key"])
    assert pre["review_status"] == "bad"


def test_v20_migration_rebuilds_old_table():
    """Пересборка v20: старые запуски целы, кандидат по умолчанию — run."""
    import os
    import sqlite3
    from migrations import migrate_to_v20
    tmp = tempfile.mkdtemp()
    con = sqlite3.connect(os.path.join(tmp, "t.sqlite"))
    con.row_factory = sqlite3.Row
    try:
        cur = con.cursor()
        cur.execute("""
            CREATE TABLE regression_runs (
                regression_id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                baseline_type TEXT NOT NULL,
                baseline_id INTEGER NOT NULL,
                candidate_run_id INTEGER NOT NULL,
                gate_max_critical INTEGER NOT NULL DEFAULT 0,
                gate_max_rate REAL NOT NULL DEFAULT 0.02,
                gate_result TEXT DEFAULT '',
                total INTEGER DEFAULT 0,
                regressions INTEGER DEFAULT 0,
                improvements INTEGER DEFAULT 0,
                unchanged INTEGER DEFAULT 0,
                created_at TEXT NOT NULL)""")
        cur.execute("""
            CREATE TABLE regression_results (
                regression_id INTEGER NOT NULL,
                stable_key TEXT NOT NULL,
                result TEXT NOT NULL,
                severity TEXT NOT NULL DEFAULT 'info',
                PRIMARY KEY (regression_id, stable_key),
                FOREIGN KEY (regression_id) REFERENCES regression_runs(regression_id))""")
        cur.execute("INSERT INTO regression_runs (name, baseline_type, baseline_id,"
                    " candidate_run_id, gate_result, total, created_at)"
                    " VALUES ('r','run',7,9,'PASS',3,'t')")
        cur.execute("INSERT INTO regression_results (regression_id, stable_key,"
                    " result) VALUES (1,'k','UNCHANGED')")
        con.commit()
        migrate_to_v20(cur)
        migrate_to_v20(cur)  # идемпотентность
        con.commit()
        cols = {r[1] for r in cur.execute("PRAGMA table_info(regression_runs)")}
        assert {"candidate_type", "candidate_version_id"} <= cols
        row = cur.execute("SELECT * FROM regression_runs").fetchone()
        assert row["candidate_type"] == "run" and row["candidate_run_id"] == 9
        assert cur.execute("SELECT COUNT(*) FROM regression_results").fetchone()[0] == 1
    finally:
        con.close()
