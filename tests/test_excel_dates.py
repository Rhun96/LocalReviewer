"""Даты Excel не роняют импорт (datetime is not JSON serializable)."""
import datetime as _dt
import tempfile

import openpyxl
from database import init_database
from file_reader import FileReader
from importer import _cell_str, import_file


def test_cell_str_dates():
    assert _cell_str(_dt.datetime(2024, 1, 15, 0, 0)) == "2024-01-15"
    assert _cell_str(_dt.datetime(2024, 1, 15, 14, 30)) == "2024-01-15 14:30"
    assert _cell_str(_dt.date(2024, 1, 15)) == "2024-01-15"
    assert _cell_str("plain") == "plain"
    assert _cell_str(42) == 42


def _xlsx_with_dates(n=1200):
    tmp = tempfile.mkdtemp()
    path = tmp + "/dates.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ID", "Дата", "Вопрос"])
    base = _dt.datetime(2024, 1, 1, 10, 30)
    for i in range(n):
        ws.append([f"D-{i}", base + _dt.timedelta(days=i), f"вопрос {i}"])
    wb.save(path)
    return path


def test_big_excel_with_datetimes_imports():
    path = _xlsx_with_dates()
    data = FileReader.read_excel_data(path, "Sheet", 0)
    assert isinstance(data[0]["Дата"], str)
    p = tempfile.mkdtemp()
    init_database(p)
    fid, n, skipped = import_file(
        p, path, "excel", "Sheet", 0,
        {"ID": "source_id", "Дата": "metadata", "Вопрос": "primary_text"},
        data)
    assert n == 1200 and skipped == 0, (n, skipped)
    from database import db
    with db(p) as conn:
        row = conn.execute("SELECT primary_text FROM cases LIMIT 1").fetchone()
        assert "вопрос" in row["primary_text"]


def test_formula_errors_become_empty_with_stats():
    """Ошибки формул (#REF! и т.п.) — пусто + счётчик, а не мусор в ревью."""
    from file_reader import is_excel_error
    assert is_excel_error("#REF!")
    assert is_excel_error("  #ссылка!  ")
    assert is_excel_error("#Н/Д")
    assert not is_excel_error("#1")
    assert not is_excel_error("#hashtag")
    assert not is_excel_error("")
    assert not is_excel_error(None)
    assert not is_excel_error(42)
    tmp = tempfile.mkdtemp()
    path = tmp + "/formulas.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ID", "Вопрос"])
    ws.append(["k1", "#REF!"])
    ws.append(["k2", "нормальный вопрос"])
    ws.append(["k3", "#ЗНАЧ!"])
    wb.save(path)
    stats: dict = {}
    data = FileReader.read_excel_data(path, "Sheet", 0, stats=stats)
    assert stats.get("formula_errors") == 2, stats
    by_id = {r["ID"]: r["Вопрос"] for r in data}
    assert by_id == {"k1": "", "k2": "нормальный вопрос", "k3": ""}
    # без stats — тоже чистим, молча
    data2 = FileReader.read_excel_data(path, "Sheet", 0)
    assert [r["Вопрос"] for r in data2] == ["", "нормальный вопрос", ""]
    # импорт тянет пустые вопросы честно (не мусор)
    p = tempfile.mkdtemp()
    init_database(p)
    fid, n, skipped = import_file(
        p, path, "excel", "Sheet", 0,
        {"ID": "source_id", "Вопрос": "primary_text"}, data)
    assert (fid, n) == (1, 3), (n, skipped)
    from database import db
    with db(p) as conn:
        texts = [r["primary_text"] for r in conn.execute(
            "SELECT primary_text FROM cases ORDER BY case_id").fetchall()]
    assert texts[0] == "" and "REF" not in "".join(texts)


def test_formula_errors_clean_in_preview_and_ods():
    """Превью маппинга и ODS — тоже без мусора формул."""
    import tempfile
    import openpyxl
    from file_reader import FileReader
    tmp = tempfile.mkdtemp()
    path = tmp + "/pv.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ID", "Вопрос"])
    ws.append(["k1", "#REF!"])
    ws.append(["k2", "ок"])
    wb.save(path)
    prev = FileReader.read_excel_preview(path, "Sheet", 100)
    assert prev["rows"][0][1] == ""
    assert prev["rows"][1][1] == "ок"
    try:
        import odf  # noqa: F401
    except ImportError:
        return
    from tests.test_ods import _make_ods
    opath = _make_ods(tmp + "/e.ods", [("S", [["q"], ["#ССЫЛКА!"], ["текст"]])])
    stats: dict = {}
    data = FileReader.read_ods_data(opath, "S", stats=stats)
    assert stats.get("formula_errors") == 1, stats
    # строка из одной ошибки пропускается целиком, как пустая
    assert [r["q"] for r in data] == ["текст"]
    prev_o = FileReader.read_ods_preview(opath, "S")
    assert prev_o["rows"][0][0] == ""
