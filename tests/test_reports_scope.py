"""Отчёты уважают скоп файла; экспорт несёт продукт и custom-колонки."""
import tempfile
from pathlib import Path

import openpyxl

from bulk_operation_service import bulk_add_tag
from database import init_database
from export_service import export_results_to_xlsx
from filter_service import get_filtered_case_ids
from importer import find_file_by_name, import_file
from report_service import get_checks_report, get_overall_report, get_tags_report
from tag_service import create_tag


def _proj():
    from autocheck_service import run_autochecks
    tmp = tempfile.mkdtemp()
    init_database(tmp)
    m = {"q": "primary_text", "a": "response_text",
         "p": "product", "c": "custom:Канал"}
    import_file(tmp, "a.xlsx", "excel", "S", 0, m,
                [{"q": "qa1", "a": "=2+2", "p": "Авиа", "c": "Чат"},
                 {"q": "qa2", "a": "ok", "p": "ЖД", "c": ""}])
    import_file(tmp, "b.xlsx", "excel", "S", 0, m,
                [{"q": "позвоните +7 (900) 111-22-33", "a": "ok",
                  "p": "Отель", "c": "Почта"}])
    run_autochecks(tmp)
    return tmp


def _fid(proj, name):
    return find_file_by_name(proj, name)["file_id"]


def test_overall_tags_checks_respect_file():
    p = _proj()
    fa, fb = _fid(p, "a.xlsx"), _fid(p, "b.xlsx")
    assert get_overall_report(p)["total"] == 3
    assert get_overall_report(p, fa)["total"] == 2
    assert get_overall_report(p, fb)["total"] == 1
    tid = create_tag(p, "T1")
    bulk_add_tag(p, get_filtered_case_ids(p, {}), tid)
    assert sum(t["cases_count"] for t in get_tags_report(p)) >= 3
    assert sum(t["cases_count"] for t in get_tags_report(p, fa)) == 2
    all_checks = get_checks_report(p)
    assert any(c["count"] > 0 for c in all_checks)
    scoped = get_checks_report(p, fa)
    assert sum(c["count"] for c in scoped) <= sum(c["count"] for c in all_checks)


def test_export_results_columns():
    p = _proj()
    out = str(Path(p) / "res.xlsx")
    assert export_results_to_xlsx(p, out) == 3
    ws = openpyxl.load_workbook(out, read_only=True).active
    headers = [c for c in next(ws.iter_rows(values_only=True))]
    assert "Продукт" in headers and "Канал" in headers
    rows = list(ws.iter_rows(values_only=True))[1:]
    by_q = {r[3]: r for r in rows}
    i_prod = headers.index("Продукт")
    i_chan = headers.index("Канал")
    assert by_q["qa1"][i_prod] == "Авиа" and by_q["qa1"][i_chan] == "Чат"
    assert by_q["qa1"][4] == "'=2+2"
